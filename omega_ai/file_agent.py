"""Omega AI v3 — File Analysis Agent
Multi-format file analysis for PDF, DOCX, XLSX, CSV, and TXT files.
"""
from __future__ import annotations

import csv
import io
import json
from pathlib import Path
from typing import Any


class FileAgent:
    """Analyze various file formats and extract meaningful content."""

    SUPPORTED = {".txt", ".csv", ".pdf", ".docx", ".xlsx", ".json", ".md"}

    def analyze(self, file_path: str) -> dict[str, Any]:
        """Analyze a file and return structured content."""
        path = Path(file_path)
        if not path.exists():
            return {"error": f"File not found: {file_path}"}

        ext = path.suffix.lower()
        if ext not in self.SUPPORTED:
            return {"error": f"Unsupported format: {ext}. Supported: {self.SUPPORTED}"}

        try:
            if ext == ".txt" or ext == ".md":
                return self._parse_text(path)
            elif ext == ".csv":
                return self._parse_csv(path)
            elif ext == ".json":
                return self._parse_json(path)
            elif ext == ".pdf":
                return self._parse_pdf(path)
            elif ext == ".docx":
                return self._parse_docx(path)
            elif ext == ".xlsx":
                return self._parse_xlsx(path)
            else:
                return {"error": f"Parser not implemented for {ext}"}
        except Exception as e:
            return {"error": f"Parse error: {type(e).__name__}: {e}"}

    def _parse_text(self, path: Path) -> dict[str, Any]:
        content = path.read_text(encoding="utf-8")
        return {"type": "text", "filename": path.name, "content": content[:5000], "lines": content.count("\n") + 1, "chars": len(content)}

    def _parse_csv(self, path: Path) -> dict[str, Any]:
        content = path.read_text(encoding="utf-8")
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        return {"type": "csv", "filename": path.name, "headers": rows[0] if rows else [], "row_count": len(rows) - 1, "preview": rows[:10], "column_count": len(rows[0]) if rows else 0}

    def _parse_json(self, path: Path) -> dict[str, Any]:
        data = json.loads(path.read_text())
        return {"type": "json", "filename": path.name, "keys": list(data.keys()) if isinstance(data, dict) else [], "item_count": len(data) if isinstance(data, (dict, list)) else 1, "preview": str(data)[:2000]}

    def _parse_pdf(self, path: Path) -> dict[str, Any]:
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            return {"type": "pdf", "filename": path.name, "pages": len(reader.pages), "content": text[:5000], "chars": len(text)}
        except ImportError:
            return {"error": "PDF parsing requires: pip install pypdf"}

    def _parse_docx(self, path: Path) -> dict[str, Any]:
        try:
            import docx
            doc = docx.Document(str(path))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            return {"type": "docx", "filename": path.name, "paragraphs": len(paragraphs), "content": "\n".join(paragraphs)[:5000], "chars": sum(len(p) for p in paragraphs)}
        except ImportError:
            return {"error": "DOCX parsing requires: pip install python-docx"}

    def _parse_xlsx(self, path: Path) -> dict[str, Any]:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(str(path), data_only=True)
            sheets_data = {}
            for sheet in wb.worksheets:
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    rows.append([str(cell) if cell is not None else "" for cell in row])
                sheets_data[sheet.title] = rows[:20]
            return {"type": "xlsx", "filename": path.name, "sheets": list(sheets_data.keys()), "sheet_count": len(wb.worksheets), "preview": sheets_data}
        except ImportError:
            return {"error": "XLSX parsing requires: pip install openpyxl"}

    def summarize(self, file_path: str) -> str:
        """Generate a human-readable summary of a file."""
        result = self.analyze(file_path)
        if "error" in result:
            return f"Error: {result['error']}"
        lines = [f"## File Analysis: {result['filename']}", f"Type: {result['type']}", ""]
        if result["type"] == "text":
            lines.append(f"Lines: {result['lines']}, Characters: {result['chars']}")
            lines.append(f"Preview:\n{result['content'][:500]}...")
        elif result["type"] == "csv":
            lines.append(f"Columns: {result['column_count']}, Rows: {result['row_count']}")
            lines.append(f"Headers: {result['headers']}")
        elif result["type"] == "json":
            lines.append(f"Keys: {result['keys']}, Items: {result['item_count']}")
        elif result["type"] == "pdf":
            lines.append(f"Pages: {result['pages']}, Characters: {result['chars']}")
        elif result["type"] == "docx":
            lines.append(f"Paragraphs: {result['paragraphs']}, Characters: {result['chars']}")
        elif result["type"] == "xlsx":
            lines.append(f"Sheets: {result['sheet_count']} - {result['sheets']}")
        return "\n".join(lines)
